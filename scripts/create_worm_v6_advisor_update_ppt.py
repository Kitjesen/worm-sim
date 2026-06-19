from __future__ import annotations

import csv
import json
import math
import os
from pathlib import Path
from typing import Iterable

import cv2
import matplotlib.pyplot as plt
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "record" / "current" / "worm_v6_advisor_update_ppt"
ASSET_DIR = OUT_DIR / "assets"

V90_SCAN_DIR = ROOT / "record" / "current" / "flat_omni_v90_server_mixed_planar_yaw_preserve_scan"
V92B_SCAN_DIR = ROOT / "record" / "current" / "flat_omni_v92b_server_low_yaw_envelope_scan"
V93B_SCAN_DIR = ROOT / "record" / "current" / "flat_omni_v93b_server_mixed_sign_repair_scan"
V94_SCAN_DIR = ROOT / "record" / "current" / "flat_omni_v94_server_robust_forward_left_scan"
V86_VIDEO_DIR = ROOT / "record" / "current" / "flat_omni_v86b_server_robust_forward_left_videos"
V85_VIDEO_DIR = ROOT / "record" / "current" / "flat_omni_v85b_server_mixed_planar_hardcase_videos"


COLORS = {
    "ink": RGBColor(25, 31, 36),
    "muted": RGBColor(93, 103, 111),
    "blue": RGBColor(41, 105, 176),
    "teal": RGBColor(18, 137, 126),
    "green": RGBColor(73, 142, 80),
    "orange": RGBColor(204, 119, 34),
    "red": RGBColor(180, 62, 54),
    "line": RGBColor(210, 216, 220),
    "panel": RGBColor(245, 247, 248),
    "white": RGBColor(255, 255, 255),
}


REFERENCE_ROWS = [
    [
        "CPG + RL",
        "Liu et al. 2020/2023",
        "稳定节律由 CPG/prior 产生，RL 调 CPG 输入或残差；支持我们 prior + residual 路线",
    ],
    [
        "contact-aware CPG",
        "Liu, Onal & Fu 2021",
        "接触/环境反馈进入节律控制；提示后续 sand/slope 不能只靠网络容量",
    ],
    [
        "snake DRL gait",
        "Shi, Dear & Kelly 2020",
        "DRL 能生成 terrestrial/aquatic snake gait，但核心是 gait generation，不是无限制全向跟踪",
    ],
    [
        "peristaltic RL",
        "Saga et al. 2016",
        "蠕动机器人可以用 Actor-Critic 学多节收缩波；对应我们的 worm prior 与轴向推进",
    ],
    [
        "sidewinding on sand",
        "Marvi et al. 2014",
        "沙地/坡地关键是接触长度和 slip 管理；后续 terrain 迁移要加入接触指标",
    ],
    [
        "velocity-command RL",
        "Margolis et al. 2024",
        "速度命令要用 curriculum 和可行速度域；高线速+高 yaw 组合可能物理不可行",
    ],
    [
        "latent gait space",
        "Mitchell et al. 2024",
        "多 gait latent space 需要 ablation 证明可解释性；对应我们的 learned gait gate",
    ],
    [
        "snake target tracking",
        "Bing et al. 2020",
        "蛇形 RL 也常把任务写成目标/路径跟踪；我们的论文应强调 body-frame command tracking",
    ],
]

REFERENCE_BULLETS = [
    "Liu et al., Learning to Locomote with Deep Neural-Network and CPG-based Control in a Soft Snake Robot, arXiv:2001.04059, doi:10.48550/arXiv.2001.04059.",
    "Liu, Onal & Fu, Reinforcement Learning of CPG-regulated Locomotion Controller for a Soft Snake Robot, IEEE T-RO / arXiv:2207.04899, doi:10.1109/TRO.2023.3286046.",
    "Liu, Onal & Fu, Learning Contact-aware CPG-based Locomotion in a Soft Snake Robot, arXiv:2105.04608.",
    "Shi, Dear & Kelly, Deep Reinforcement Learning for Snake Robot Locomotion, IFAC-PapersOnLine 2020, doi:10.1016/j.ifacol.2020.12.2619.",
    "Saga et al., Acquisition of earthworm-like movement patterns of many-segmented peristaltic crawling robots, International Journal of Advanced Robotic Systems 2016, doi:10.1177/1729881416657740.",
    "Marvi et al., Sidewinding with minimal slip: Snake and robot ascent of sandy slopes, Science 2014, doi:10.1126/science.1255718.",
    "Margolis et al., Rapid locomotion via reinforcement learning, IJRR 2024, doi:10.1177/02783649231224053.",
    "Mitchell et al., Gaitor: Learning a Unified Representation Across Gaits for Real-World Quadruped Locomotion, arXiv:2405.19452.",
    "Bing et al., Perception-Action Coupling Target Tracking Control for a Snake Robot via Reinforcement Learning, Frontiers in Neurorobotics 2020, doi:10.3389/fnbot.2020.591128.",
]

REFERENCE_LINKS = [
    "https://arxiv.org/abs/2001.04059",
    "https://arxiv.org/abs/2207.04899",
    "https://arxiv.org/abs/2105.04608",
    "https://www.sciencedirect.com/science/article/pii/S2405896320333772",
    "https://journals.sagepub.com/doi/abs/10.1177/1729881416657740",
    "https://arxiv.org/abs/1410.2945",
    "https://journals.sagepub.com/doi/10.1177/02783649231224053",
    "https://arxiv.org/abs/2405.19452",
    "https://www.frontiersin.org/journals/neurorobotics/articles/10.3389/fnbot.2020.591128/full",
]


def add_textbox(slide, text: str, x: float, y: float, w: float, h: float, size: int = 18,
                color=COLORS["ink"], bold: bool = False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title(slide, title: str, subtitle: str | None = None):
    add_textbox(slide, title, 0.55, 0.32, 12.1, 0.55, size=24, bold=True)
    if subtitle:
        add_textbox(slide, subtitle, 0.58, 0.86, 11.7, 0.33, size=10, color=COLORS["muted"])
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.18), Inches(12.2), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = COLORS["line"]
    line.line.color.rgb = COLORS["line"]


def add_panel(slide, x: float, y: float, w: float, h: float, fill=COLORS["panel"]):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = COLORS["line"]
    return shp


def add_bullets(slide, bullets: Iterable[str], x: float, y: float, w: float, h: float,
                size: int = 15, color=COLORS["ink"], gap: float = 0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(gap * 72)
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = color
    return box


def load_v90_summary() -> list[dict]:
    p = V90_SCAN_DIR / "acceptance_summary.json"
    return json.loads(p.read_text(encoding="utf-8"))["items"]


def load_optional_acceptance(scan_dir: Path) -> list[dict]:
    p = scan_dir / "acceptance_summary.json"
    if not p.exists():
        return []
    return json.loads(p.read_text(encoding="utf-8")).get("items", [])


def v92b_status_bullets() -> list[str]:
    rows = load_optional_acceptance(V92B_SCAN_DIR)
    if not rows:
        return [
            "V92b strict scan 尚未拉取/汇总；当前正式基线仍为 V90。",
        ]
    accepted = {row["label"]: bool(row.get("accepted")) for row in rows}
    final_robust = next(
        (row for row in rows if row["label"] == "v92b_final_robust"), None)
    bullets = [
        "V92b 已完成 strict scan：final_nominal/final_robust accepted，但 best_robust 有 1 个 planar sign fail。",
        "结论：V92b final 是下一轮视频/ablation 候选；V90 仍是最稳的 all-checkpoint 基线。",
    ]
    if final_robust:
        bullets.append(
            "V92b final_robust: planar RMSE "
            f"{final_robust['planar_rmse_m_s']:.3f} m/s, yaw RMSE "
            f"{final_robust['yaw_rmse_rad_s']:.3f} rad/s, planar/off-axis "
            f"exceed {final_robust['planar_error_exceed_count']}/"
            f"{final_robust['off_axis_exceed_count']}."
        )
    if accepted.get("v92b_best_robust") is False:
        bullets.append(
            "V92b best_robust counterexample: cmd=(0.05, 0.075, 0), measured vx<0，说明混合命令仍需修。"
        )
    return bullets


def v93b_status_bullets() -> list[str]:
    rows = load_optional_acceptance(V93B_SCAN_DIR)
    if not rows:
        return [
            "V93b strict scan 尚未拉取/汇总；下一步仍是验证 mixed vx/vy 鲁棒错号。",
        ]
    accepted = {row["label"]: bool(row.get("accepted")) for row in rows}
    final_robust = next(
        (row for row in rows if row["label"] == "v93b_final_robust"), None)
    bullets = [
        "V93b 已完成：best/final 的 nominal 均 accepted；best/final robust 均失败。",
        "失败条件非常集中：robust 下 cmd=(0.05, 0.075, 0) 出现负向 vx 投影，仍属 mixed_vx_vy counterexample。",
    ]
    if final_robust:
        bullets.append(
            "V93b final_robust: planar RMSE "
            f"{final_robust['planar_rmse_m_s']:.3f} m/s, yaw RMSE "
            f"{final_robust['yaw_rmse_rad_s']:.3f} rad/s, wrong sign "
            f"{final_robust['wrong_planar_sign_count']}/"
            f"{final_robust['wrong_yaw_sign_count']}."
        )
    if accepted.get("v93b_best_robust") is False:
        bullets.append(
            "结论：V93b 不能替代 V90/V92b；下一轮要针对 mixed forward-left robust sign 做课程或先验修正。"
        )
    return bullets


def v94_status_bullets() -> list[str]:
    rows = load_optional_acceptance(V94_SCAN_DIR)
    if not rows:
        return [
            "V94 robust-forward-left targeted scan 尚未拉取/汇总。",
        ]
    final_robust = next(
        (row for row in rows if row["label"] == "v94_final_robust"), None)
    bullets = [
        "V94 已完成 targeted hardcase 训练和 scan：nominal 过，robust best/final 仍失败。",
        "结论：继续同类 hardcase 课程不足以修复 forward-left mixed sign，需要改先验/奖励结构或做分阶段残差约束。",
    ]
    if final_robust:
        bullets.append(
            "V94 final_robust: planar RMSE "
            f"{final_robust['planar_rmse_m_s']:.3f} m/s, yaw RMSE "
            f"{final_robust['yaw_rmse_rad_s']:.3f} rad/s, selected cmd "
            f"({final_robust['counter_cmd_vx_m_s']:.2f}, "
            f"{final_robust['counter_cmd_vy_m_s']:.2f}, "
            f"{final_robust['counter_cmd_yaw_rad_s']:.1f})."
        )
    return bullets


def load_video_manifest(path: Path) -> list[dict]:
    p = path / "video_manifest.md"
    rows = []
    if not p.exists():
        return rows
    for line in p.read_text(encoding="utf-8").splitlines():
        if not line.startswith("| ") or "---" in line or "Case" in line:
            continue
        parts = [x.strip() for x in line.strip("|").split("|")]
        if len(parts) < 6:
            continue
        rows.append({
            "case": parts[0],
            "duration": parts[1],
            "frames": parts[2],
            "speed": parts[3],
            "yaw": parts[4],
            "gate": parts[5],
            "mp4": parts[6].split("(")[-1].rstrip(")") if len(parts) > 6 else "",
        })
    return rows


def extract_frame(video: Path, out: Path, label: str, at_ratio: float = 0.36) -> bool:
    cap = cv2.VideoCapture(str(video))
    if not cap.isOpened():
        return False
    n = int(cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
    idx = max(0, min(n - 1, int(n * at_ratio))) if n else 0
    cap.set(cv2.CAP_PROP_POS_FRAMES, idx)
    ok, frame = cap.read()
    cap.release()
    if not ok:
        return False
    frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
    img = Image.fromarray(frame)
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("arial.ttf", 42)
    except Exception:
        font = ImageFont.load_default()
    draw.rectangle([20, 20, 20 + 24 * len(label), 76], fill=(255, 255, 255))
    draw.text((34, 30), label, fill=(20, 30, 40), font=font)
    img.save(out)
    return True


def make_metric_chart(items: list[dict]) -> Path:
    labels = [x["label"].replace("v90_", "").replace("_", "\n") for x in items]
    planar = [x["planar_rmse_m_s"] for x in items]
    yaw = [x["yaw_rmse_rad_s"] for x in items]
    x = range(len(labels))
    fig, ax1 = plt.subplots(figsize=(10, 4.8), dpi=180)
    ax1.bar([i - 0.18 for i in x], planar, width=0.36, label="planar RMSE (m/s)", color="#2b6cb0")
    ax1.axhline(0.10, color="#2f855a", linestyle="--", linewidth=1.2, label="planar gate 0.10")
    ax1.set_ylabel("planar RMSE (m/s)")
    ax1.set_ylim(0, 0.11)
    ax1.set_xticks(list(x))
    ax1.set_xticklabels(labels, fontsize=8)
    ax2 = ax1.twinx()
    ax2.bar([i + 0.18 for i in x], yaw, width=0.36, label="yaw RMSE (rad/s)", color="#dd6b20")
    ax2.axhline(0.20, color="#c05621", linestyle=":", linewidth=1.2, label="yaw gate 0.20")
    ax2.set_ylabel("yaw RMSE (rad/s)")
    ax2.set_ylim(0, 0.22)
    lines, names = ax1.get_legend_handles_labels()
    lines2, names2 = ax2.get_legend_handles_labels()
    ax1.legend(lines + lines2, names + names2, loc="upper right", fontsize=8)
    ax1.set_title("V90 35-command strict scan: nominal and robust all accepted")
    ax1.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    out = ASSET_DIR / "v90_metric_chart.png"
    fig.savefig(out)
    plt.close(fig)
    return out


def make_contact_sheet(thumbnails: list[tuple[str, Path]], out: Path) -> Path:
    cell_w, cell_h = 520, 330
    cols = 3
    rows = math.ceil(len(thumbnails) / cols)
    sheet = Image.new("RGB", (cols * cell_w, rows * cell_h), "white")
    draw = ImageDraw.Draw(sheet)
    try:
        font = ImageFont.truetype("arial.ttf", 28)
    except Exception:
        font = ImageFont.load_default()
    for idx, (label, img_path) in enumerate(thumbnails):
        r, c = divmod(idx, cols)
        x, y = c * cell_w, r * cell_h
        img = Image.open(img_path).convert("RGB")
        img.thumbnail((cell_w - 24, cell_h - 72))
        sheet.paste(img, (x + 12, y + 52))
        draw.text((x + 18, y + 14), label, fill=(25, 31, 36), font=font)
        draw.rectangle([x, y, x + cell_w - 1, y + cell_h - 1], outline=(210, 216, 220), width=2)
    sheet.save(out)
    return out


def make_assets() -> dict[str, Path | list[tuple[str, Path]]]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    items = load_v90_summary()
    metric_chart = make_metric_chart(items)
    cases = [
        ("forward", "forward_12s_1080p.mp4"),
        ("reverse", "reverse_12s_1080p.mp4"),
        ("lateral_left", "lateral_left_12s_1080p.mp4"),
        ("lateral_right", "lateral_right_12s_1080p.mp4"),
        ("yaw_left", "yaw_left_12s_1080p.mp4"),
        ("yaw_right", "yaw_right_12s_1080p.mp4"),
        ("continuous_sweep", "continuous_sweep_24s_1080p.mp4"),
        ("hard_forward_left", "hard_forward_left_12s_1080p.mp4"),
        ("hard_forward_right", "hard_forward_right_12s_1080p.mp4"),
    ]
    thumbs = []
    for label, name in cases:
        video = V86_VIDEO_DIR / name
        if not video.exists():
            continue
        out = ASSET_DIR / f"thumb_{label}.png"
        if extract_frame(video, out, label):
            thumbs.append((label, out))
    sheet = make_contact_sheet(thumbs, ASSET_DIR / "video_contact_sheet.png")
    return {"metric_chart": metric_chart, "thumbs": thumbs, "contact_sheet": sheet}


def add_image(slide, path: Path, x: float, y: float, w: float, h: float | None = None):
    if h is None:
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    else:
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def add_table(slide, rows: list[list[str]], x: float, y: float, w: float, h: float,
              font_size: int = 9, header: bool = True):
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for c in range(len(rows[0])):
        table.columns[c].width = Inches(w / len(rows[0]))
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(font_size)
            p.font.color.rgb = COLORS["ink"]
            if header and r == 0:
                p.font.bold = True
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(231, 238, 245)
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLORS["white"]
    return table


def build_deck() -> Path:
    assets = make_assets()
    v90 = load_v90_summary()
    v92b_notes = v92b_status_bullets()
    v93b_notes = v93b_status_bullets()
    v94_notes = v94_status_bullets()
    videos = load_video_manifest(V86_VIDEO_DIR)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor(238, 242, 244)
    add_textbox(s, "Worm V6 双模态全向运动训练进展", 0.65, 0.72, 11.6, 0.75, size=30, bold=True)
    add_textbox(s, "蛇形 + 蠕虫 gait prior，PPO residual，latent gait gate，实机可部署观测契约", 0.68, 1.48, 11.3, 0.35, size=15, color=COLORS["muted"])
    add_panel(s, 0.72, 2.18, 5.7, 2.2, fill=COLORS["white"])
    add_bullets(s, [
        "当前最强证据：V90 flat 35-command nominal/robust strict scan 全部 accepted",
        "接口保持实机兼容：80D obs，12D action = 11D residual + 1D gait gate",
        "仍不宣称无限制全向：mixed vx/vy 幅值误差与串扰仍需继续训练",
    ], 1.0, 2.52, 5.1, 1.55, size=15)
    add_image(s, assets["contact_sheet"], 6.75, 2.0, 5.7, 3.65)
    add_textbox(s, "V94 targeted 加训仍未修复 robust mixed vx/vy 错号；V90/V92b 仍是当前汇报基线", 0.75, 6.65, 9.8, 0.32, size=12, color=COLORS["blue"], bold=True)

    # 2
    s = prs.slides.add_slide(blank)
    add_title(s, "这次汇报先讲清楚边界：不是麦克纳姆轮式“任意全向”", "蛇/蠕虫机器人靠身体-地面摩擦推进，vx、vy、yaw 不是独立轮式驱动通道")
    add_panel(s, 0.65, 1.55, 3.85, 4.85)
    add_textbox(s, "目标", 0.9, 1.82, 2.7, 0.35, size=19, bold=True, color=COLORS["blue"])
    add_bullets(s, [
        "在可行速度包络内跟踪 body-frame vx, vy, yaw_rate",
        "让蛇形与蠕虫步态都参与，而不是固定单一 gait",
        "观测量必须能从真实机器人获得",
    ], 0.92, 2.25, 3.2, 2.5, size=14)
    add_panel(s, 4.85, 1.55, 3.85, 4.85)
    add_textbox(s, "当前最好说法", 5.1, 1.82, 3.1, 0.35, size=19, bold=True, color=COLORS["teal"])
    add_bullets(s, [
        "finite-envelope command-conditioned multimodal planar locomotion",
        "flat 地形上通过 35-command nominal/robust strict scan",
        "不是所有连续速度都已高精度跟踪",
    ], 5.12, 2.25, 3.25, 2.55, size=14)
    add_panel(s, 9.05, 1.55, 3.05, 4.85)
    add_textbox(s, "剩余风险", 9.3, 1.82, 2.4, 0.35, size=19, bold=True, color=COLORS["orange"])
    add_bullets(s, [
        "mixed vx/vy 幅值误差",
        "off-axis crosstalk",
        "纯 yaw 会有平移漂移风险",
        "sand/slope 尚未正式迁移",
    ], 9.32, 2.25, 2.4, 2.8, size=14)

    # 3
    s = prs.slides.add_slide(blank)
    add_title(s, "机器人接口按实机部署设计：每节 encoder + IMU")
    add_panel(s, 0.65, 1.55, 5.8, 4.95)
    add_textbox(s, "80D Observation", 0.95, 1.82, 4.8, 0.32, size=20, bold=True, color=COLORS["blue"])
    rows = [
        ["Group", "Dim", "Source"],
        ["command", "3", "vx, vy, yaw_rate"],
        ["joint_pos", "11", "6 slide + 5 yaw encoders"],
        ["joint_vel", "11", "encoder velocity estimate"],
        ["previous_action", "11", "last deployed motor command"],
        ["segment_gravity", "21", "7 IMU projected gravity"],
        ["segment_gyro", "21", "7 IMU angular velocity"],
        ["phase_clock", "2", "sin/cos, 1s period"],
    ]
    add_table(s, rows, 0.88, 2.35, 5.25, 3.45, font_size=8)
    add_panel(s, 6.75, 1.55, 5.75, 4.95)
    add_textbox(s, "12D Action", 7.05, 1.82, 4.8, 0.32, size=20, bold=True, color=COLORS["teal"])
    add_bullets(s, [
        "11D residual motor action：在 gait prior 上做修正",
        "1D learned latent gait gate z_g：连续混合 peristaltic 与 serpentine priors",
        "部署动作：prior + scaled residual + EMA/saturation",
        "actor / critic: MLP 512-256-128, SB3 ActorCriticPolicy",
    ], 7.05, 2.35, 4.95, 2.5, size=14)
    add_textbox(s, "论文可写：The policy outputs an 11-dimensional residual motor action and a one-dimensional learned latent gait gate z_g.", 7.05, 5.45, 4.95, 0.55, size=10, color=COLORS["muted"])

    # 4
    s = prs.slides.add_slide(blank)
    add_title(s, "控制结构：gait prior 给稳定节律，PPO 学残差与门控")
    add_panel(s, 0.7, 1.55, 12.0, 4.85, fill=COLORS["white"])
    steps = [
        ("Command", "vx, vy, yaw"),
        ("Policy", "80D obs -> 12D action"),
        ("Gate", "z_g blends worm/snake priors"),
        ("Prior", "CMA-ES rhythmic motor primitives"),
        ("Adapter", "residual + EMA + saturation"),
        ("Robot", "6 slide + 5 yaw joints"),
    ]
    x0 = 1.0
    for i, (h, b) in enumerate(steps):
        x = x0 + i * 1.95
        add_panel(s, x, 2.45, 1.55, 1.15, fill=RGBColor(234, 243, 247))
        add_textbox(s, h, x + 0.12, 2.62, 1.3, 0.25, size=13, bold=True, align=PP_ALIGN.CENTER, color=COLORS["blue"])
        add_textbox(s, b, x + 0.08, 2.92, 1.38, 0.4, size=8, align=PP_ALIGN.CENTER, color=COLORS["muted"])
        if i < len(steps) - 1:
            add_textbox(s, "→", x + 1.58, 2.77, 0.45, 0.3, size=22, bold=True, color=COLORS["muted"])
    add_bullets(s, [
        "和文献里的 CPG + RL 思路一致：不让网络从零裸学全部电机波形。",
        "V50 componentwise mixed-command composition 被实验否定，当前默认是 hardcase-gated directional adapter。",
        "V93b 从 V92b final 继续后 nominal 通过，但 robust mixed_vx_vy 仍有 1 个方向错号。",
    ], 1.0, 4.55, 10.9, 1.0, size=13)

    # 5
    s = prs.slides.add_slide(blank)
    add_title(s, "最新定量结果：V90 是当前最强 flat 验收候选")
    add_image(s, assets["metric_chart"], 0.75, 1.45, 6.8, 3.25)
    rows = [["Case", "Accepted", "Planar RMSE", "Yaw RMSE", "Wrong sign", "Exceed"]]
    for item in v90:
        rows.append([
            item["label"].replace("v90_", ""),
            "yes" if item["accepted"] else "no",
            f'{item["planar_rmse_m_s"]:.3f}',
            f'{item["yaw_rmse_rad_s"]:.3f}',
            f'{item["wrong_planar_sign_count"]}/{item["wrong_yaw_sign_count"]}',
            f'planar {item["planar_error_exceed_count"]}, off {item["off_axis_exceed_count"]}',
        ])
    add_table(s, rows, 7.8, 1.52, 4.9, 2.75, font_size=8)
    add_bullets(s, [
        "V90 best/final 的 nominal 与 robust 35-command strict analysis 全部 accepted。",
        "robust 条件包含 encoder/IMU noise、1 step action delay、0.90 action saturation。",
        "限制：accepted 不代表每个命令都无误差；仍有 planar exceed / off-axis exceed。",
    ], 0.9, 5.15, 11.6, 1.1, size=13)
    add_bullets(s, v92b_notes[:2] + v93b_notes[:1] + v94_notes[:2], 7.85, 4.30, 4.7, 1.45, size=7, color=COLORS["muted"])

    # 6
    s = prs.slides.add_slide(blank)
    add_title(s, "视频展示页：六方向 + 连续速度变化 + hard diagonal")
    add_image(s, assets["contact_sheet"], 0.72, 1.42, 8.2, 5.2)
    rows = [["Video", "Speed", "Yaw", "Gate"]]
    for row in videos[:9]:
        rows.append([row["case"].replace("_12s_1080p", "").replace("_24s_1080p", ""), row["speed"], row["yaw"], row["gate"]])
    add_table(s, rows, 9.25, 1.45, 3.05, 4.45, font_size=7)
    add_textbox(s, f"视频目录：{V86_VIDEO_DIR}", 0.75, 6.77, 11.3, 0.25, size=8, color=COLORS["muted"])

    # 7
    s = prs.slides.add_slide(blank)
    add_title(s, "轨迹和时程图：看的是每个体节的绝对轨迹，不再全部从 0 起画")
    add_image(s, V86_VIDEO_DIR / "continuous_sweep_24s_1080p_trajectory.png", 0.65, 1.45, 6.25, 4.45)
    add_image(s, V86_VIDEO_DIR / "continuous_sweep_24s_1080p_telemetry.png", 7.0, 1.45, 5.8, 4.45)
    add_bullets(s, [
        "左：head + segment 1-6 的世界坐标 XY 轨迹、body-frame forward/lateral/Z 时程。",
        "右：命令速度 vs 实测速度、learned gate、prior/residual/applied action heatmap。",
    ], 0.8, 6.18, 11.6, 0.65, size=12)

    # 8
    s = prs.slides.add_slide(blank)
    add_title(s, "训练版本演进：问题从“方向不分”推进到“混合命令误差/串扰”")
    rows = [
        ["Stage", "做了什么", "主要结论"],
        ["V13-V29", "vx/vy/yaw 初始课程 + speed gate", "能动但 yaw 与 lateral 弱，误差大"],
        ["V35-V41", "轴向/横向/yaw 分离修复", "yaw drift 降低，lateral 仍不稳"],
        ["V50-V59", "mixed command composition / hardcase", "V50 组合路径被否定，需 gated prior"],
        ["V70-V77", "服务器 robust hardcase gate", "重复 sign failure 被压住"],
        ["V85b-V90", "mixed-planar + yaw-preserve 修复", "V90 通过 nominal/robust strict scan"],
        ["V91-V92b", "继续训练，低 yaw 包络", "V91 退化；V92b final accepted，但 best_robust 暴露 mixed sign fail"],
    ]
    add_table(s, rows, 0.78, 1.45, 11.8, 4.25, font_size=10)
    add_bullets(s, [
        "现在不再只看 PPO reward；模型选择按 strict scan、方向符号、off-axis、zero drift、yaw-only 漂移。",
        "服务器训练为正式路径；本机只做脚本、文档、轻量验证和 PPT。",
    ], 0.9, 6.12, 11.2, 0.75, size=12)

    # 9
    s = prs.slides.add_slide(blank)
    add_title(s, "机械建模：真实 slide 是单向拉绳收缩 + 钢片被动回弹")
    add_panel(s, 0.75, 1.5, 5.6, 4.95)
    add_textbox(s, "当前 MuJoCo 训练模型", 1.0, 1.82, 4.5, 0.35, size=18, bold=True, color=COLORS["blue"])
    add_bullets(s, [
        "全身训练使用等效 slide spring：约 300 N/m，阻尼 10",
        "视觉钢片由 box 段渲染，随体节姿态更新，不参与受力",
        "优点：稳定、快，适合 PPO",
    ], 1.0, 2.35, 4.8, 2.0, size=14)
    add_panel(s, 6.75, 1.5, 5.6, 4.95)
    add_textbox(s, "论文/实机下一步", 7.0, 1.82, 4.5, 0.35, size=18, bold=True, color=COLORS["teal"])
    add_bullets(s, [
        "单节高保真 cable-strip 标定：力-位移曲线拟合 k_eq",
        "真实机构用 unilateral cable/tendon：电机只拉不推",
        "用拉伸试验机采集钢片力-位移和动态释放数据",
    ], 7.0, 2.35, 4.9, 2.0, size=14)

    # 10
    s = prs.slides.add_slide(blank)
    add_title(s, "文献对比：我们不是从零裸学电机，而是 prior + residual + gate")
    rows = [["Literature line", "代表工作", "对我们的启发"]] + REFERENCE_ROWS
    add_table(s, rows, 0.55, 1.25, 12.2, 4.95, font_size=7)
    add_textbox(s, "结论：我们的创新定位应是“可部署观测契约 + 双模态 gait prior + 学习 gate + 有限包络速度跟踪”。", 0.75, 6.12, 11.3, 0.45, size=14, bold=True, color=COLORS["blue"])

    # 11
    s = prs.slides.add_slide(blank)
    add_title(s, "论文创新点：先守住能证明的贡献")
    add_panel(s, 0.72, 1.45, 3.8, 4.95)
    add_textbox(s, "C1 实机观测契约", 1.0, 1.78, 3.0, 0.32, size=17, bold=True, color=COLORS["blue"])
    add_bullets(s, ["80D obs 全来自命令、编码器、IMU、上一动作和 phase clock", "避免使用仿真不可部署的全局状态"], 1.0, 2.35, 3.0, 1.8, size=13)
    add_panel(s, 4.75, 1.45, 3.8, 4.95)
    add_textbox(s, "C2 双模态 gait gate", 5.03, 1.78, 3.0, 0.32, size=17, bold=True, color=COLORS["teal"])
    add_bullets(s, ["11D residual + 1D z_g", "连续混合 peristaltic / serpentine priors", "后续做 fixed-gate / no-gate ablation"], 5.03, 2.35, 3.0, 2.2, size=13)
    add_panel(s, 8.78, 1.45, 3.8, 4.95)
    add_textbox(s, "C3 有限包络全向", 9.06, 1.78, 3.0, 0.32, size=17, bold=True, color=COLORS["orange"])
    add_bullets(s, ["flat 35-command nominal/robust strict scan 已有 accepted 证据", "继续推进 mixed command 精度、sand/slope 迁移和实机验证"], 9.06, 2.35, 3.0, 2.0, size=13)

    # 12
    s = prs.slides.add_slide(blank)
    add_title(s, "当前问题：训练不是没学会，而是物理包络和混合命令仍难")
    add_bullets(s, [
        "V90 通过 strict scan，但 mixed vx/vy 和 forward-yaw 仍会出现 planar exceed / off-axis exceed。",
        "V91 延续同一课程后退化，说明继续堆训练步数不一定带来更好策略。",
        "纯 yaw 命令容易诱发平移漂移或身体收缩成团；V92b/V93b 低 yaw 包络已把 yaw RMSE 降低到约 0.02 rad/s。",
        "V93b 证明当前剩余核心不是 yaw RMSE，而是 robust mixed vx/vy 的方向符号。",
        "V94 targeted robust-forward-left 加训仍失败，说明下一步不能只靠同类课程继续堆步数。",
        "V50 componentwise mixed-command composition 不是默认方向，实验显示会增加 planar sign failure。",
        "正式论文还需要 fixed-gate/no-gate ablation、能耗/稳定性、sand/slope、实机数据。",
    ], 0.9, 1.55, 11.6, 4.3, size=16)
    add_textbox(s, "训练策略：不要扩大命令域，先把可行包络内的方向符号、幅值、串扰和 yaw-only 稳定性打稳。", 0.9, 6.18, 11.0, 0.42, size=14, bold=True, color=COLORS["blue"])

    # 13
    s = prs.slides.add_slide(blank)
    add_title(s, "下一步计划：从 flat 可信结果走向论文闭环")
    rows = [
        ["Priority", "Action", "Acceptance evidence"],
        ["1", "用 V90 与 V92b final 录制正式高清视频", "forward/reverse/lateral/yaw + continuous sweep"],
        ["2", "改 mixed sign repair 机制 / gate ablation", "V92b-V94 都指向 cmd=(0.05,0.075,0) 反向 vx"],
        ["3", "gate ablation", "fixed worm/mixed/snake, no-gate, learned gate, gate sweep"],
        ["4", "钢片/单向拉绳参数采集", "力-位移、阶跃响应、IMU、摩擦、质量几何 CSV"],
        ["5", "sand/slope 迁移", "flat 通过后再正式启动，增加 slip/contact 指标"],
        ["6", "实机部署预检", "ABI fingerprint, deploy bundle, hardware log/video"],
    ]
    add_table(s, rows, 0.62, 1.45, 12.1, 4.8, font_size=10)

    # 14
    s = prs.slides.add_slide(blank)
    add_title(s, "导师讨论点")
    add_bullets(s, [
        "论文命名：建议用 finite-envelope command-conditioned multimodal planar locomotion，而不是直接 full omnidirectional。",
        "是否接受“钢片高保真单节标定 + 全身等效弹簧训练”的两级建模路线？",
        "目标期刊/会议要求决定实验深度：只做 flat/sand/slope 仿真，还是必须做实机三地形？",
        "gate 是否作为核心贡献，需要 ablation 证明；否则只作为工程模块描述。",
        "下一轮是否优先硬件参数采集，以减少 sim-to-real 质疑？",
    ], 0.9, 1.55, 11.6, 4.55, size=17)
    add_textbox(s, "附录/素材：PPT 输出目录包含视频链接清单、V90 scan 数据和文献来源列表。", 0.9, 6.45, 10.8, 0.3, size=11, color=COLORS["muted"])

    # 15
    s = prs.slides.add_slide(blank)
    add_title(s, "References used for this update")
    add_bullets(s, REFERENCE_BULLETS, 0.8, 1.35, 12.0, 5.7, size=9)

    out = OUT_DIR / "worm_v6_advisor_update_20260603.pptx"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


def write_manifest(pptx_path: Path):
    manifest = {
        "pptx": str(pptx_path),
        "created_from": {
            "v90_scan_dir": str(V90_SCAN_DIR),
            "v92b_scan_dir": str(V92B_SCAN_DIR),
            "v93b_scan_dir": str(V93B_SCAN_DIR),
            "v94_scan_dir": str(V94_SCAN_DIR),
            "video_dir": str(V86_VIDEO_DIR),
            "fallback_video_dir": str(V85_VIDEO_DIR),
        },
        "server_training_status": {
            "v91": "completed; callback metrics regressed relative to V90",
            "v92b": (
                "completed; final_nominal/final_robust accepted; "
                "best_robust failed one planar sign command"
            ),
            "v93b": (
                "completed from V92b final; nominal best/final accepted; "
                "robust best/final failed one mixed_vx_vy planar sign command"
            ),
            "v94": (
                "completed from V92b final with robust sensor/delay/saturation "
                "training; nominal accepted but robust best/final still failed "
                "one mixed_vx_vy planar sign command"
            ),
        },
        "limitations": [
            "PPT uses V90 as the safest all-checkpoint strict baseline.",
            "V92b final is a candidate improvement, but V92b best_robust failed one planar sign command.",
            "V93b did not supersede V90/V92b because robust scans still fail one mixed_vx_vy sign command.",
            "V94 targeted hardcase training also did not supersede V90/V92b.",
            "PPT uses V86b videos because V90/V92b formal videos are not yet recorded locally.",
            "Video files are linked by local path rather than embedded to keep deck size manageable.",
        ],
        "presentation_runtime_note": (
            "@oai/artifact-tool/presentation-jsx is unavailable in this "
            "workspace, so this deck is generated as an editable python-pptx "
            "PowerPoint rather than a raster-only deck."
        ),
        "reference_links": REFERENCE_LINKS,
    }
    (OUT_DIR / "ppt_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    links = [
        "# Worm V6 advisor update PPT manifest",
        "",
        f"- PPTX: `{pptx_path}`",
        f"- V90 scan: `{V90_SCAN_DIR}`",
        f"- V92b scan: `{V92B_SCAN_DIR}`",
        f"- V93b scan: `{V93B_SCAN_DIR}`",
        f"- V94 scan: `{V94_SCAN_DIR}`",
        f"- Video directory: `{V86_VIDEO_DIR}`",
        "- Runtime note: generated as editable python-pptx because artifact-tool presentation-jsx is unavailable here.",
        "",
        "## Key video files",
    ]
    for name in [
        "forward_12s_1080p.mp4",
        "reverse_12s_1080p.mp4",
        "lateral_left_12s_1080p.mp4",
        "lateral_right_12s_1080p.mp4",
        "yaw_left_12s_1080p.mp4",
        "yaw_right_12s_1080p.mp4",
        "continuous_sweep_24s_1080p.mp4",
    ]:
        links.append(f"- `{V86_VIDEO_DIR / name}`")
    links.extend(["", "## Literature source links"])
    for url in REFERENCE_LINKS:
        links.append(f"- {url}")
    (OUT_DIR / "ppt_manifest.md").write_text("\n".join(links) + "\n", encoding="utf-8")


def validate_pptx(path: Path):
    prs = Presentation(path)
    assert len(prs.slides) == 15, len(prs.slides)
    assert path.stat().st_size > 100_000, path.stat().st_size
    print(f"pptx={path}")
    print(f"slides={len(prs.slides)}")
    print(f"size_bytes={path.stat().st_size}")


def main():
    pptx_path = build_deck()
    write_manifest(pptx_path)
    validate_pptx(pptx_path)


if __name__ == "__main__":
    main()
