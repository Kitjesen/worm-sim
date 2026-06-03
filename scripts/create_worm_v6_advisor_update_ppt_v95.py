from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

import create_worm_v6_advisor_update_ppt as base


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "record" / "current" / "worm_v6_advisor_update_ppt"
V95_RUN_LABEL = "flat_random_v95_server_mixed_positive_vx_hardcase_from_v92bfinal_np2"
V95_RUN_DIR = (
    "/home/bsrl/hongsenpang/codex_runs/worm-sim-v6-server-training/runs/"
    f"worm_v6_ppo_{V95_RUN_LABEL}"
)
V95_LOG_DIR = (
    "/home/bsrl/hongsenpang/codex_runs/worm-sim-v6-server-training/"
    "server_logs"
)
V95_SCAN_DIR = (
    ROOT / "record" / "current"
    / "flat_omni_v95_server_mixed_positive_vx_hardcase_scan"
)

V95_COMPONENT_AUDIT_ROWS = [
    [
        "best nominal",
        "false",
        "0.0547",
        "0.0217",
        "7",
        "(+0.05,+0.075,0)->(-0.0229,+0.0578,+0.1136)",
    ],
    [
        "best robust",
        "false",
        "0.0570",
        "0.0201",
        "8",
        "(+0.05,+0.075,0)->(-0.0345,+0.0059,+0.0113)",
    ],
    [
        "final nominal",
        "false",
        "0.0572",
        "0.0218",
        "11",
        "(+0.05,+0.075,0)->(-0.0267,+0.0516,+0.0851)",
    ],
    [
        "final robust",
        "false",
        "0.0616",
        "0.0201",
        "10",
        "(-0.10,+0.075,0)->(-0.0605,-0.0338,-0.0882)",
    ],
]


PAPER_ROWS = [
    [
        "Liu et al. 2023",
        "Soft snake CPG-RL",
        "RL outputs CPG tonic input; CPG keeps rhythmic form",
        "Use prior/CPG to constrain exploration; RL learns modulation",
    ],
    [
        "Liu et al. 2021/2023",
        "Contact-aware soft snake",
        "Goal controller plus contact/reflex controller around CPG",
        "Sand/slope should add contact/slip feedback, not only larger MLP",
    ],
    [
        "Shi et al. 2020",
        "Snake DRL gait discovery",
        "Direct DRL can discover gaits in simulation",
        "Useful ablation: end-to-end RL is harder than prior + residual",
    ],
    [
        "Bing et al. 2020",
        "Snake target tracking",
        "PPO controls joint positions for visual target tracking",
        "Good for task reward design, but less deployable gait structure",
    ],
    [
        "Liu et al. 2022",
        "Camera path following",
        "Hierarchical RL outputs gait offsets over a gait equation",
        "Closest path-tracking analogy: RL modifies gait parameters",
    ],
    [
        "Sartoretti et al. 2019",
        "Decentralized articulated robots",
        "Distributed local policies for articulated locomotion",
        "Future option for many segments: shared local segment policy",
    ],
    [
        "Bellegarda & Ijspeert 2022",
        "CPG-RL velocity tracking",
        "RL modulates oscillator setpoints for omnidirectional commands",
        "Strong support for command-conditioned prior modulation",
    ],
    [
        "Kim et al. 2021",
        "Hierarchical multi-gait RL",
        "Velocity commands induce multiple gait regimes",
        "Supports learned latent gait gate and fixed-gate ablation",
    ],
]


def add_v95_status_slide(prs: Presentation) -> None:
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    base.add_title(
        slide,
        "V95 已启动：针对 robust mixed forward-left 错号做局部奖励修复",
        "本轮不是简单增加步数，而是把 V92b-V94 重复失败的慢速 +vx/full-lateral 命令写进奖励契约。",
    )
    base.add_panel(slide, 0.65, 1.45, 5.95, 4.95)
    base.add_textbox(
        slide,
        "训练设置",
        0.95,
        1.78,
        4.6,
        0.35,
        size=18,
        bold=True,
        color=base.COLORS["blue"],
    )
    base.add_bullets(
        slide,
        [
            "服务器训练已启动，run label: V95 mixed-positive-vx hardcase",
            "从 V92b final 继续，不从 V94 继续，避免继承 targeted 退化",
            "actor/critic 都是 512-256-128，保持 80D obs / 12D action ABI",
            "鲁棒条件：encoder/IMU noise + 1-step action delay + 0.90 saturation",
            "训练块：100k steps；完成后跑 nominal/robust strict scan",
        ],
        0.95,
        2.3,
        5.1,
        2.7,
        size=13,
    )

    base.add_panel(slide, 6.95, 1.45, 5.65, 4.95)
    base.add_textbox(
        slide,
        "V95 契约新增项",
        7.25,
        1.78,
        4.6,
        0.35,
        size=18,
        bold=True,
        color=base.COLORS["teal"],
    )
    rows = [
        ["Item", "Value"],
        ["target commands", "(+0.05,+/-0.075,0)"],
        ["gate", "+vx, full lateral, yaw=0"],
        ["penalty", "forward deficit below positive target"],
        ["weight", "6.0"],
        ["ABI", "unchanged"],
    ]
    base.add_table(slide, rows, 7.25, 2.3, 4.9, 2.1, font_size=9)
    base.add_bullets(
        slide,
        [
            "目的：把“总 RMSE 看起来还可以，但 forward 分量已经反向”的情况显式惩罚。",
            "边界：V95 结果出来前，汇报仍以 V90/V92b 作为已验证结果。",
        ],
        7.25,
        4.65,
        4.9,
        0.95,
        size=11,
        color=base.COLORS["muted"],
    )
    base.add_textbox(
        slide,
        f"Server run: {V95_RUN_DIR}",
        0.75,
        6.72,
        11.8,
        0.25,
        size=8,
        color=base.COLORS["muted"],
    )


def add_v95_component_audit_slide(prs: Presentation) -> None:
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    base.add_title(
        slide,
        "V95 结果审计：投影方向改善，但 mixed vx/vy 分量仍未学会",
        "旧 analyzer 会把 final robust 判为 accepted；新增分量符号 gate 后，V95 仍失败。",
    )
    base.add_panel(slide, 0.55, 1.18, 12.25, 4.55)
    rows = [[
        "checkpoint",
        "accepted",
        "planar RMSE",
        "yaw RMSE",
        "wrong mixed comp",
        "representative component failure",
    ]] + V95_COMPONENT_AUDIT_ROWS
    base.add_table(slide, rows, 0.85, 1.55, 11.65, 2.75, font_size=7)
    base.add_bullets(
        slide,
        [
            "V95 的真实进步：final checkpoint 在旧投影方向 gate 下 robust 扫描不再有 wrong planar/yaw。",
            "Metric keyword: mixed-component sign gate = wrong_mixed_component_sign_count。",
            "新的问题定义：mixed vx/vy 命令必须同时满足 vx、vy 两个非零分量同号，不能只看投影方向。",
            "结论边界：V95 是诊断性进步，不是连续 vx/vy/yaw 跟踪完成版。",
            "下一步训练：奖励与选择标准直接加入 mixed-component sign 和 component magnitude，不再只优化 planar projection。",
        ],
        0.85,
        4.55,
        11.2,
        1.05,
        size=12,
    )
    base.add_textbox(
        slide,
        f"Scan artifacts: {V95_SCAN_DIR}",
        0.75,
        6.72,
        11.8,
        0.25,
        size=8,
        color=base.COLORS["muted"],
    )


def add_literature_training_slide(prs: Presentation) -> None:
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    base.add_title(
        slide,
        "已有论文怎么训练：主流路线是 gait/CPG prior + RL 调制",
        "这支持我们当前的 prior + residual + learned gait gate，而不是端到端裸输出所有电机。",
    )
    rows = [["Paper", "Task", "Control form", "Takeaway"]] + PAPER_ROWS
    base.add_table(slide, rows, 0.45, 1.25, 12.45, 5.15, font_size=6)
    base.add_textbox(
        slide,
        "汇报结论：我们的目标应写成有限速度包络内的 command-conditioned multimodal locomotion；全向要给可行速度域和严格 scan 约束。",
        0.75,
        6.62,
        11.5,
        0.35,
        size=12,
        bold=True,
        color=base.COLORS["blue"],
    )


def add_video_index_slide(prs: Presentation) -> None:
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    base.add_title(
        slide,
        "视频和图表素材索引：导师汇报时可按这个顺序播放",
        "视频文件未嵌入 PPT，使用本地绝对路径/清单管理，避免 PPT 过大。",
    )
    rows = [
        ["Purpose", "File"],
        ["六方向 3x2 总览", "record/current/flat_omni_v86b_server_robust_forward_left_videos"],
        ["forward / reverse", "forward_12s_1080p.mp4 / reverse_12s_1080p.mp4"],
        ["lateral left/right", "lateral_left_12s_1080p.mp4 / lateral_right_12s_1080p.mp4"],
        ["yaw left/right", "yaw_left_12s_1080p.mp4 / yaw_right_12s_1080p.mp4"],
        ["连续速度变化", "continuous_sweep_24s_1080p.mp4"],
        ["轨迹图", "continuous_sweep_24s_1080p_trajectory.png"],
        ["速度/gait/action 遥测", "continuous_sweep_24s_1080p_telemetry.png"],
        ["V90 strict scan", "record/current/flat_omni_v90_server_mixed_planar_yaw_preserve_scan"],
        ["V94 failed hardcase", "record/current/flat_omni_v94_server_robust_forward_left_scan"],
    ]
    base.add_table(slide, rows, 0.7, 1.35, 11.9, 4.65, font_size=9)
    base.add_bullets(
        slide,
        [
            "播放顺序建议：六方向总览 -> continuous sweep -> trajectory/telemetry -> V90 scan -> V94 failure -> V95 plan。",
            "展示边界：这些视频主要是 V86b 视觉材料；严格定量结论以 V90/V92b/V94 scan JSON 为准。",
        ],
        0.9,
        6.18,
        11.2,
        0.75,
        size=12,
    )


def write_v95_manifest(pptx_path: Path, slide_count: int) -> None:
    manifest = {
        "pptx": str(pptx_path),
        "slides": slide_count,
        "created_from_base_script": str(ROOT / "scripts" / "create_worm_v6_advisor_update_ppt.py"),
        "v95_server_training": {
            "status": "completed_not_accepted_under_component_sign_gate",
            "run_label": V95_RUN_LABEL,
            "run_dir": V95_RUN_DIR,
            "log_dir": V95_LOG_DIR,
            "resume": "V92b final",
            "curriculum": "robust_forward_left_diagonal_repair",
            "train_chunk_steps": 100000,
            "reward_change": "mixed_positive_vx_full_lateral_deficit_penalty",
        },
        "v95_component_audit": {
            "scan_dir": str(V95_SCAN_DIR),
            "strict_component_gate": "wrong_mixed_component_sign_count",
            "summary": "All four V95 scans fail once mixed vx/vy component sign is required.",
            "rows": V95_COMPONENT_AUDIT_ROWS,
        },
        "video_dir": str(base.V86_VIDEO_DIR),
        "scan_dirs": {
            "v90": str(base.V90_SCAN_DIR),
            "v92b": str(base.V92B_SCAN_DIR),
            "v93b": str(base.V93B_SCAN_DIR),
            "v94": str(base.V94_SCAN_DIR),
        },
        "literature_rows": PAPER_ROWS,
        "limitations": [
            "V95 completed, but it is not accepted under the stricter mixed-component sign gate.",
            "After completion, V95 final robust passes the old projection gate but fails the stricter mixed-component sign gate.",
            "V90 remains the safest fully accepted flat strict-scan baseline.",
            "V86b videos are visual evidence; formal claims use scan JSON/CSV.",
            "The deck is editable python-pptx because presentation-jsx is unavailable in this workspace.",
        ],
    }
    (OUT_DIR / "ppt_manifest_v95.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    lines = [
        "# Worm V6 advisor update PPT V95 manifest",
        "",
        f"- PPTX: `{pptx_path}`",
        f"- Slides: `{slide_count}`",
        f"- V95 status: completed, but not accepted under mixed-component sign gate; run label `{V95_RUN_LABEL}`",
        f"- V95 run dir: `{V95_RUN_DIR}`",
        f"- V95 scan dir: `{V95_SCAN_DIR}`",
        f"- Video dir: `{base.V86_VIDEO_DIR}`",
        "",
        "## Key point",
        "",
        "V95 is a diagnostic improvement: final robust passes the old projected-direction gate, but all four scans fail the stricter mixed-component sign gate.",
    ]
    (OUT_DIR / "ppt_manifest_v95.md").write_text(
        "\n".join(lines) + "\n",
        encoding="utf-8",
    )


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    base_path = base.build_deck()
    prs = Presentation(base_path)
    add_v95_status_slide(prs)
    add_v95_component_audit_slide(prs)
    add_literature_training_slide(prs)
    add_video_index_slide(prs)
    out = OUT_DIR / "worm_v6_advisor_update_20260603_v95.pptx"
    prs.save(out)
    slide_count = len(Presentation(out).slides)
    assert slide_count == 19, slide_count
    assert out.stat().st_size > 100_000, out.stat().st_size
    write_v95_manifest(out, slide_count)
    print(f"pptx={out}")
    print(f"slides={slide_count}")
    print(f"size_bytes={out.stat().st_size}")


if __name__ == "__main__":
    main()
